"""
API des documents ComptaFlow.

Fonctions ajoutées:
- suppression définitive;
- saisie pour toutes les catégories;
- validation des factures et relevés bancaires;
- modification des mouvements bancaires;
- export CSV et Excel des données extraites.
"""

import csv
import hashlib
import io
import json
import logging
import mimetypes
import re
import uuid
from datetime import datetime, timezone

from pathlib import Path
from typing import Any

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Query,
    UploadFile,
)
from fastapi.responses import (
    FileResponse,
    Response,
)

from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel, Field

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import get_db
from app.core.deps import (
    get_current_user,
    require_role,
)

from app.models.document import Document
from app.models.ecriture import EcritureComptable
from app.models.entreprise import Entreprise
from app.models.enums import (
    RoleEnum,
    StatutDocumentEnum,
    StatutValidationEnum,
)
from app.models.mouvement_bancaire import (
    MouvementBancaire,
)
from app.models.user import User

from app.schemas.document import DocumentEntrepriseUpdate, DocumentOut
from app.schemas.document_detail import (
    DocumentDetailOut,
)
from app.schemas.mouvement_bancaire import (
    MouvementBancaireOut,
    MouvementBancaireUpdate,
)

from app.tasks.document_processing import (
    process_document,
)
from app.services import audit_service, rapprochement_bancaire_service, workflow_comptable_service


router = APIRouter(
    prefix="/documents",
    tags=["documents"],
)


logger = logging.getLogger(__name__)


ALLOWED_MIME_TYPES = {
    "application/pdf",
    "image/png",
    "image/jpeg",
    "image/jpg",
}


MAX_FILE_SIZE_MB = 20
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024


class DocumentUploadBatchAudit(BaseModel):
    document_ids: list[uuid.UUID] = Field(min_length=2, max_length=100)

_EXTENSIONS_PAR_MIME = {
    "application/pdf": {".pdf"},
    "image/png": {".png"},
    "image/jpeg": {".jpg", ".jpeg"},
    "image/jpg": {".jpg", ".jpeg"},
}


def _detecter_mime_reel(contenu: bytes) -> str | None:
    if contenu.startswith(b"%PDF-"):
        return "application/pdf"
    if contenu.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if contenu.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    return None


def _valider_fichier_uploade(
    contenu: bytes,
    nom_fichier: str,
    mime_declare: str | None,
) -> tuple[str, str]:
    if len(contenu) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(status_code=400, detail="Fichier trop volumineux (maximum 20 MB).")

    mime_reel = _detecter_mime_reel(contenu)
    if mime_reel is None:
        raise HTTPException(status_code=400, detail="Contenu de fichier non reconnu (PDF, PNG ou JPEG requis).")

    if mime_declare not in ALLOWED_MIME_TYPES:
        raise HTTPException(status_code=400, detail="Type MIME déclaré non autorisé.")

    mime_declare_normalise = "image/jpeg" if mime_declare == "image/jpg" else mime_declare
    if mime_declare_normalise != mime_reel:
        raise HTTPException(status_code=400, detail="Le contenu du fichier ne correspond pas au type MIME déclaré.")

    extension = Path(nom_fichier).suffix.lower()
    if extension not in _EXTENSIONS_PAR_MIME[mime_reel]:
        raise HTTPException(status_code=400, detail="L'extension ne correspond pas au contenu du fichier.")

    return mime_reel, extension


def _resoudre_chemin_stockage(chemin: str) -> Path:
    racine = Path(settings.STORAGE_PATH).resolve()
    candidat = Path(chemin).resolve()
    if not candidat.is_relative_to(racine):
        raise HTTPException(status_code=404, detail="Fichier original introuvable sur le serveur.")
    return candidat


_ROLES_VALIDATION = (
    RoleEnum.ADMIN_CABINET,
    RoleEnum.EXPERT_COMPTABLE,
    RoleEnum.CHEF_MISSION,
)


def _get_document_or_404(
    db: Session,
    document_id: uuid.UUID,
    cabinet_id: uuid.UUID,
) -> Document:
    document = db.execute(
        select(
            Document
        ).where(
            Document.id == document_id,
            Document.cabinet_id == cabinet_id,
        )
    ).scalar_one_or_none()

    if document is None:
        raise HTTPException(
            status_code=404,
            detail="Document introuvable.",
        )

    return document


def _get_first_entry(
    db: Session,
    document_id: uuid.UUID,
) -> EcritureComptable | None:
    return db.execute(
        select(
            EcritureComptable
        ).where(
            EcritureComptable.document_id
            == document_id
        )
    ).scalars().first()


def _get_movements(
    db: Session,
    document_id: uuid.UUID,
) -> list[MouvementBancaire]:
    return list(
        db.execute(
            select(
                MouvementBancaire
            )
            .where(
                MouvementBancaire.document_id
                == document_id
            )
            .order_by(
                MouvementBancaire.date_operation,
                MouvementBancaire.created_at,
            )
        ).scalars().all()
    )


def _build_document_detail(
    db: Session,
    document: Document,
) -> DocumentDetailOut:
    return DocumentDetailOut(
        id=document.id,

        nom_fichier_original=(
            document.nom_fichier_original
        ),

        taille_octets=(
            document.taille_octets
        ),

        mime_type=document.mime_type,

        statut=document.statut,

        created_at=document.created_at,

        entreprise_id=(
            document.entreprise_id
        ),

        annee=document.annee,
        mois=document.mois,

        categorie=(
            document.categorie.value
            if document.categorie
            else None
        ),

        texte_ocr=document.texte_ocr,

        donnees_extraites=(
            document.donnees_extraites
        ),

        message_erreur=(
            document.message_erreur
        ),

        type_erreur=(
            document.type_erreur.value
            if document.type_erreur
            else None
        ),

        error_code=document.error_code,

        saisie_topaze=(
            document.saisie_topaze
        ),

        ecriture=_get_first_entry(
            db,
            document.id,
        ),

        mouvements_bancaires=(
            _get_movements(
                db,
                document.id,
            )
        ),
    )


def _enum_value(
    value: Any,
) -> Any:
    return (
        value.value
        if hasattr(value, "value")
        else value
    )


def _safe_export_value(
    value: Any,
) -> str:
    if value is None:
        return ""

    if isinstance(
        value,
        (
            dict,
            list,
            tuple,
        ),
    ):
        text = json.dumps(
            value,
            ensure_ascii=False,
            default=str,
        )
    else:
        text = str(
            _enum_value(value)
        )

    # Protection contre l'injection de formules Excel.
    if text.startswith(
        (
            "=",
            "+",
            "-",
            "@",
        )
    ):
        return "'" + text

    return text


def _flatten_json(
    value: Any,
    prefix: str = "",
) -> dict[str, str]:
    result: dict[str, str] = {}

    if isinstance(
        value,
        dict,
    ):
        for key, child in value.items():
            new_key = (
                f"{prefix}.{key}"
                if prefix
                else str(key)
            )

            result.update(
                _flatten_json(
                    child,
                    new_key,
                )
            )

        return result

    if isinstance(
        value,
        list,
    ):
        result[
            prefix or "valeur"
        ] = _safe_export_value(
            value
        )

        return result

    result[
        prefix or "valeur"
    ] = _safe_export_value(
        value
    )

    return result


def _document_metadata(
    document: Document,
) -> dict[str, str]:
    return {
        "document_id": str(
            document.id
        ),

        "fichier": (
            document.nom_fichier_original
        ),

        "categorie": (
            _safe_export_value(
                document.categorie
            )
        ),

        "annee": (
            _safe_export_value(
                document.annee
            )
        ),

        "mois": (
            _safe_export_value(
                document.mois
            )
        ),

        "statut_document": (
            _safe_export_value(
                document.statut
            )
        ),

        "saisie_topaze": (
            "Oui"
            if document.saisie_topaze
            else "Non"
        ),

        "date_import": (
            document.created_at.isoformat()
        ),
    }


def _entry_data(
    entry: EcritureComptable | None,
) -> dict[str, str]:
    if entry is None:
        return {}

    return {
        "ecriture_id": str(
            entry.id
        ),

        "type_ecriture": (
            _safe_export_value(
                entry.type_ecriture
            )
        ),

        "numero_piece": (
            _safe_export_value(
                entry.numero_piece
            )
        ),

        "date_piece": (
            _safe_export_value(
                entry.date_piece
            )
        ),

        "tiers": (
            _safe_export_value(
                entry.tiers
            )
        ),

        "montant_ht": (
            _safe_export_value(
                entry.montant_ht
            )
        ),

        "taux_tva": (
            _safe_export_value(
                entry.taux_tva
            )
        ),

        "montant_tva": (
            _safe_export_value(
                entry.montant_tva
            )
        ),

        "montant_ttc": (
            _safe_export_value(
                entry.montant_ttc
            )
        ),

        "statut_validation": (
            _safe_export_value(
                entry.statut_validation
            )
        ),

        "anomalie_detectee": (
            "Oui"
            if entry.anomalie_detectee
            else "Non"
        ),

        "anomalie_details": (
            _safe_export_value(
                entry.anomalie_details
            )
        ),
    }


def _safe_filename(
    value: str,
) -> str:
    stem = Path(
        value
    ).stem

    cleaned = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        stem,
    ).strip("_")

    return (
        cleaned[:80]
        or "document"
    )


def _generate_csv_export(
    document: Document,
    entry: EcritureComptable | None,
    movements: list[
        MouvementBancaire
    ],
) -> bytes:
    stream = io.StringIO()

    writer = csv.writer(
        stream,
        delimiter=";",
    )

    writer.writerow(
        ["DONNÉES DU DOCUMENT"]
    )

    writer.writerow(
        [
            "Champ",
            "Valeur",
        ]
    )

    flattened = _flatten_json(
        document.donnees_extraites
        or {}
    )

    combined = {
        **_document_metadata(
            document
        ),

        **_entry_data(
            entry
        ),

        **{
            f"extraction.{key}": value
            for key, value
            in flattened.items()
            if not key.startswith(
                "lignes_bancaires"
            )
        },
    }

    for key, value in combined.items():
        writer.writerow(
            [
                key,
                value,
            ]
        )

    if movements:
        writer.writerow([])
        writer.writerow(
            ["MOUVEMENTS BANCAIRES"]
        )

        writer.writerow(
            [
                "Date",
                "Libellé",
                "Référence",
                "Type",
                "Montant",
                "Solde après opération",
            ]
        )

        for movement in movements:
            writer.writerow(
                [
                    (
                        movement
                        .date_operation
                        .isoformat()
                    ),

                    _safe_export_value(
                        movement.libelle
                    ),

                    _safe_export_value(
                        movement.reference
                    ),

                    _safe_export_value(
                        movement.type_mouvement
                    ),

                    _safe_export_value(
                        movement.montant
                    ),

                    _safe_export_value(
                        movement
                        .solde_apres_operation
                    ),
                ]
            )

    return (
        "\ufeff"
        + stream.getvalue()
    ).encode("utf-8")


def _generate_xlsx_export(
    document: Document,
    entry: EcritureComptable | None,
    movements: list[
        MouvementBancaire
    ],
) -> bytes:
    workbook = Workbook()

    data_sheet = workbook.active
    data_sheet.title = (
        "Données extraites"
    )

    data_sheet.append(
        [
            "Champ",
            "Valeur",
        ]
    )

    for cell in data_sheet[1]:
        cell.font = Font(
            bold=True
        )

    flattened = _flatten_json(
        document.donnees_extraites
        or {}
    )

    combined = {
        **_document_metadata(
            document
        ),

        **_entry_data(
            entry
        ),

        **{
            f"extraction.{key}": value
            for key, value
            in flattened.items()
            if not key.startswith(
                "lignes_bancaires"
            )
        },
    }

    for key, value in combined.items():
        data_sheet.append(
            [
                key,
                value,
            ]
        )

    data_sheet.column_dimensions[
        "A"
    ].width = 35

    data_sheet.column_dimensions[
        "B"
    ].width = 65

    if movements:
        movement_sheet = (
            workbook.create_sheet(
                "Mouvements bancaires"
            )
        )

        movement_sheet.append(
            [
                "Date",
                "Libellé",
                "Référence",
                "Type",
                "Montant",
                "Solde après opération",
            ]
        )

        for cell in movement_sheet[1]:
            cell.font = Font(
                bold=True
            )

        for movement in movements:
            movement_sheet.append(
                [
                    (
                        movement
                        .date_operation
                        .isoformat()
                    ),

                    _safe_export_value(
                        movement.libelle
                    ),

                    _safe_export_value(
                        movement.reference
                    ),

                    _safe_export_value(
                        movement.type_mouvement
                    ),

                    float(
                        movement.montant
                    ),

                    (
                        float(
                            movement
                            .solde_apres_operation
                        )
                        if movement
                        .solde_apres_operation
                        is not None
                        else None
                    ),
                ]
            )

        widths = [
            14,
            45,
            20,
            14,
            16,
            22,
        ]

        for column, width in zip(
            "ABCDEF",
            widths,
        ):
            movement_sheet.column_dimensions[
                column
            ].width = width

    stream = io.BytesIO()

    workbook.save(
        stream
    )

    stream.seek(0)

    return stream.read()


def _delete_physical_file(
    file_path_value: str,
) -> bool:
    """
    Supprime uniquement un fichier qui se trouve réellement
    dans le dossier STORAGE_PATH.
    """
    file_path = Path(
        file_path_value
    ).resolve()

    storage_root = Path(
        settings.STORAGE_PATH
    ).resolve()

    try:
        file_path.relative_to(
            storage_root
        )

    except ValueError:
        return False

    try:
        if (
            file_path.exists()
            and file_path.is_file()
        ):
            file_path.unlink()

        return True

    except OSError:
        return False


def _delete_accounting_data(
    db: Session,
    document_id: uuid.UUID,
) -> None:
    """
    Supprime les données comptables liées à un document.

    Avant de supprimer les écritures, on détache les références
    ``doublon_potentiel_id`` provenant d'autres écritures. Cette colonne
    est une clé étrangère auto-référencée et peut bloquer la suppression
    avec une violation de contrainte PostgreSQL.
    """
    entry_ids = list(
        db.execute(
            select(EcritureComptable.id).where(
                EcritureComptable.document_id == document_id
            )
        ).scalars().all()
    )

    if entry_ids:
        db.query(EcritureComptable).filter(
            EcritureComptable.doublon_potentiel_id.in_(entry_ids)
        ).update(
            {EcritureComptable.doublon_potentiel_id: None},
            synchronize_session=False,
        )

    db.query(MouvementBancaire).filter(
        MouvementBancaire.document_id == document_id
    ).delete(synchronize_session=False)

    db.query(EcritureComptable).filter(
        EcritureComptable.document_id == document_id
    ).delete(synchronize_session=False)


@router.post(
    "/upload",
    response_model=DocumentOut,
)
def upload_document(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    content = file.file.read(MAX_FILE_SIZE_BYTES + 1)

    file_hash = hashlib.sha256(
        content
    ).hexdigest()

    existing = db.execute(
        select(
            Document
        ).where(
            Document.cabinet_id
            == current_user.cabinet_id,

            Document.hash_fichier
            == file_hash,
        )
    ).scalar_one_or_none()

    if existing is not None:
        raise HTTPException(
            status_code=409,
            detail=(
                "Ce fichier a déjà été uploadé "
                f"(document {existing.id})."
            ),
        )

    storage_dir = (
        Path(
            settings.STORAGE_PATH
        )
        / str(
            current_user.cabinet_id
        )
    )

    storage_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    original_filename = (
        file.filename
        or "document"
    )

    mime_reel, file_extension = _valider_fichier_uploade(
        content, original_filename, file.content_type
    )

    stored_path = (
        storage_dir
        / (
            f"{uuid.uuid4()}"
            f"{file_extension}"
        )
    )

    with stored_path.open(
        "wb"
    ) as target:
        target.write(
            content
        )

    document = Document(
        cabinet_id=(
            current_user.cabinet_id
        ),

        uploaded_by=(
            current_user.id
        ),

        nom_fichier_original=(
            original_filename
        ),

        chemin_stockage=str(
            stored_path
        ),

        hash_fichier=file_hash,

        taille_octets=len(
            content
        ),

        mime_type=mime_reel,

        saisie_topaze=False,
    )

    try:
        db.add(document)
        db.flush()
        audit_service.enregistrer(
            db, user=current_user,
            action=audit_service.AuditAction.DOCUMENTS_UPLOADED,
            resource_type="document", resource_id=document.id,
            description=f"Import du document {original_filename}.",
            metadata={"filenames": [original_filename], "mime_type": mime_reel},
            item_count=1, resource_ids=[document.id],
            correlation_id=str(uuid.uuid4()),
        )
        db.commit()
        db.refresh(document)

    except Exception:
        db.rollback()

        stored_path.unlink(
            missing_ok=True
        )

        raise

    process_document.delay(
        str(document.id)
    )

    return document


@router.post("/upload/batch-audit", status_code=204)
def audit_upload_batch(
    payload: DocumentUploadBatchAudit,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    documents = db.execute(select(Document).where(
        Document.cabinet_id == current_user.cabinet_id,
        Document.id.in_(payload.document_ids),
    )).scalars().all()
    if len(documents) != len(set(payload.document_ids)):
        raise HTTPException(status_code=404, detail="Un document importé est introuvable dans ce cabinet.")
    correlation_id = str(uuid.uuid4())
    audit_service.enregistrer(
        db, user=current_user, action=audit_service.AuditAction.DOCUMENTS_UPLOADED,
        resource_type="document", description=f"Import groupé de {len(documents)} documents.",
        metadata={
            "filenames": [item.nom_fichier_original for item in documents],
            "invoice_numbers": [
                (item.donnees_extraites or {}).get("numero_piece")
                for item in documents if isinstance(item.donnees_extraites, dict)
            ],
        },
        item_count=len(documents), resource_ids=[item.id for item in documents],
        correlation_id=correlation_id,
    )
    db.commit()
    return Response(status_code=204)


@router.get(
    "",
    response_model=list[DocumentOut],
)
def list_documents(
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    return list(
        db.execute(
            select(
                Document
            )
            .where(
                Document.cabinet_id
                == current_user.cabinet_id
            )
            .order_by(
                Document.created_at.desc()
            )
        ).scalars().all()
    )


@router.get(
    "/{document_id}",
    response_model=DocumentDetailOut,
)
def get_document_detail(
    document_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    document = _get_document_or_404(
        db,
        document_id,
        current_user.cabinet_id,
    )

    return _build_document_detail(
        db,
        document,
    )


@router.get(
    "/{document_id}/fichier"
)
def telecharger_fichier_original(
    document_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    document = _get_document_or_404(
        db,
        document_id,
        current_user.cabinet_id,
    )

    file_path = _resoudre_chemin_stockage(document.chemin_stockage)

    if not file_path.exists():
        raise HTTPException(
            status_code=404,
            detail=(
                "Fichier original introuvable "
                "sur le serveur."
            ),
        )

    media_type = (
        document.mime_type
        or mimetypes.guess_type(
            str(file_path)
        )[0]
        or "application/octet-stream"
    )

    audit_service.enregistrer(
        db, user=current_user, action=audit_service.AuditAction.DOCUMENT_OPENED,
        entreprise_id=document.entreprise_id, resource_type="document",
        resource_id=document.id,
        description=f"Consultation du fichier original {document.nom_fichier_original}.",
    )
    db.commit()

    return FileResponse(
        path=file_path,

        media_type=media_type,

        filename=Path(document.nom_fichier_original).name.replace("\r", "").replace("\n", ""),
        content_disposition_type="inline",
    )


@router.patch(
    "/{document_id}/saisie"
)
def toggle_document_saisie(
    document_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    document = _get_document_or_404(
        db,
        document_id,
        current_user.cabinet_id,
    )
    workflow_comptable_service.verifier_document_modifiable(db, document)

    ancien_statut_topaze = bool(document.saisie_topaze)
    target = not document.saisie_topaze
    entries = db.execute(
        select(
            EcritureComptable
        ).where(
            EcritureComptable.document_id
            == document.id
        )
    ).scalars().all()

    for entry in entries:
        try:
            workflow_comptable_service.verifier_periode_modifiable(db, entry)
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        if target and entry.statut_validation not in {
            StatutValidationEnum.PRETE_TOPAZE,
            StatutValidationEnum.SAISIE_TOPAZE,
            StatutValidationEnum.VALIDE,
        }:
            raise HTTPException(status_code=422, detail="Le document comporte une pré-écriture qui n'est pas prête pour Topaze.")
        entry.saisie_topaze = target
        entry.statut_validation = (
            StatutValidationEnum.SAISIE_TOPAZE if target else StatutValidationEnum.PRETE_TOPAZE
        )
        entry.ready_for_topaze_at = entry.ready_for_topaze_at or datetime.now(timezone.utc)
        entry.topaze_entered_at = datetime.now(timezone.utc) if target else None
        entry.topaze_entered_by = current_user.id if target else None
    document.saisie_topaze = target

    audit_service.enregistrer(
        db, user=current_user,
        action=audit_service.AuditAction.DOCUMENT_MARKED_TOPAZE,
        entreprise_id=document.entreprise_id, resource_type="document",
        resource_id=document.id,
        description="Modification du statut de saisie Topaze du document.",
        avant={"saisie_topaze": ancien_statut_topaze},
        apres={"saisie_topaze": bool(document.saisie_topaze)},
    )
    db.commit()

    return {
        "document_id": str(
            document.id
        ),

        "saisie_topaze": (
            document.saisie_topaze
        ),
    }


@router.patch(
    "/{document_id}/validate",
    response_model=DocumentDetailOut,
)
def validate_document(
    document_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(
            *_ROLES_VALIDATION
        )
    ),
):
    document = _get_document_or_404(
        db,
        document_id,
        current_user.cabinet_id,
    )
    workflow_comptable_service.verifier_document_modifiable(db, document)

    entry = _get_first_entry(
        db,
        document.id,
    )

    ancien_statut = document.statut.value if hasattr(document.statut, "value") else str(document.statut)
    if entry is not None:
        try:
            workflow_comptable_service.verifier_periode_modifiable(db, entry)
            anomalies = workflow_comptable_service.controler_et_transitionner(
                db, entry, document=document, user=current_user, actor_type="user"
            )
        except ValueError as exc:
            db.rollback()
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        if anomalies:
            db.commit()
            raise HTTPException(status_code=422, detail="Contrôles non satisfaits : " + " | ".join(item["message"] for item in anomalies))
    else:
        document.statut = StatutDocumentEnum.VALIDE

    audit_service.enregistrer(
        db, user=current_user, action=audit_service.AuditAction.DOCUMENT_VALIDATED,
        entreprise_id=document.entreprise_id,
        resource_type="document", resource_id=document.id,
        description="Relance des contrôles du document et de sa pré-écriture associée.",
        avant={"statut": ancien_statut}, apres={"statut": StatutDocumentEnum.VALIDE.value},
    )

    db.commit()
    db.refresh(document)

    return _build_document_detail(
        db,
        document,
    )


@router.patch(
    "/{document_id}/reject",
    response_model=DocumentDetailOut,
)
def reject_document(
    document_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(
            *_ROLES_VALIDATION
        )
    ),
):
    document = _get_document_or_404(
        db,
        document_id,
        current_user.cabinet_id,
    )
    workflow_comptable_service.verifier_document_modifiable(db, document)

    entry = _get_first_entry(
        db,
        document.id,
    )

    ancien_statut = document.statut.value if hasattr(document.statut, "value") else str(document.statut)
    if entry is not None:
        try:
            workflow_comptable_service.verifier_periode_modifiable(db, entry)
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
        entry.statut_validation = (
            StatutValidationEnum.REJETE
        )

        entry.validated_by = (
            current_user.id
        )

    # Le document reste traité,
    # mais non validé.
    document.statut = (
        StatutDocumentEnum.TRAITE
    )

    audit_service.enregistrer(
        db, user=current_user, action=audit_service.AuditAction.DOCUMENT_REJECTED,
        entreprise_id=document.entreprise_id,
        resource_type="document", resource_id=document.id,
        description="Rejet du document pour vérification ou correction.",
        avant={"statut": ancien_statut}, apres={"statut": StatutDocumentEnum.TRAITE.value},
    )

    db.commit()
    db.refresh(document)

    return _build_document_detail(
        db,
        document,
    )


@router.patch(
    "/{document_id}/mouvements/{movement_id}",
    response_model=MouvementBancaireOut,
)
def update_bank_movement(
    document_id: uuid.UUID,
    movement_id: uuid.UUID,
    payload: MouvementBancaireUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(
            *_ROLES_VALIDATION
        )
    ),
):
    document = _get_document_or_404(
        db,
        document_id,
        current_user.cabinet_id,
    )

    movement = db.execute(
        select(
            MouvementBancaire
        ).where(
            MouvementBancaire.id
            == movement_id,

            MouvementBancaire.document_id
            == document.id,

            MouvementBancaire.cabinet_id
            == current_user.cabinet_id,
        )
    ).scalar_one_or_none()

    if movement is None:
        raise HTTPException(
            status_code=404,
            detail=(
                "Mouvement bancaire "
                "introuvable."
            ),
        )

    data = payload.model_dump(
        exclude_unset=True
    )
    workflow_comptable_service.verifier_mouvement_modifiable(
        db,
        movement,
        nouvelle_date=data.get("date_operation"),
    )
    avant_mouvement = {
        key: getattr(movement, key) for key in data
    }

    required_fields = {
        "date_operation",
        "libelle",
        "type_mouvement",
        "montant",
    }

    for (
        field_name,
        field_value,
    ) in data.items():
        if (
            field_name
            in required_fields
            and field_value is None
        ):
            raise HTTPException(
                status_code=422,
                detail=(
                    f"Le champ {field_name} "
                    "ne peut pas être vide."
                ),
            )

        setattr(
            movement,
            field_name,
            field_value,
        )

    # Toute correction bancaire invalide le rapprochement précédent puis
    # relance immédiatement la proposition avec les nouvelles valeurs.
    rapprochement_bancaire_service.annuler_rapprochement(
        db,
        movement
    )
    db.flush()
    rapprochement_bancaire_service.rapprocher_mouvement(
        db,
        movement,
    )

    document.statut = (
        StatutDocumentEnum.TRAITE
    )

    apres_mouvement = {key: getattr(movement, key) for key in data}
    avant_modifie, apres_modifie = audit_service.valeurs_modifiees(
        avant_mouvement, apres_mouvement
    )
    audit_service.enregistrer(
        db, user=current_user,
        action=audit_service.AuditAction.BANK_MOVEMENT_UPDATED,
        entreprise_id=movement.entreprise_id,
        resource_type="mouvement_bancaire", resource_id=movement.id,
        description="Modification d'un mouvement bancaire extrait.",
        avant=avant_modifie, apres=apres_modifie,
    )
    db.commit()
    db.refresh(movement)

    return movement


@router.get(
    "/{document_id}/export/{format}"
)
def export_document_data(
    document_id: uuid.UUID,
    format: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    normalized_format = (
        format.lower()
    )

    if normalized_format not in {
        "csv",
        "xlsx",
    }:
        raise HTTPException(
            status_code=400,
            detail=(
                "Format non supporté. "
                "Formats acceptés: csv, xlsx."
            ),
        )

    document = _get_document_or_404(
        db,
        document_id,
        current_user.cabinet_id,
    )

    entry = _get_first_entry(
        db,
        document.id,
    )

    movements = _get_movements(
        db,
        document.id,
    )

    base_name = _safe_filename(
        document.nom_fichier_original
    )

    if normalized_format == "csv":
        content = _generate_csv_export(
            document,
            entry,
            movements,
        )

        media_type = (
            "text/csv; charset=utf-8"
        )

        filename = (
            "donnees_extraites_"
            f"{base_name}.csv"
        )

    else:
        content = _generate_xlsx_export(
            document,
            entry,
            movements,
        )

        media_type = (
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )

        filename = (
            "donnees_extraites_"
            f"{base_name}.xlsx"
        )

    audit_service.enregistrer(
        db, user=current_user, action=audit_service.AuditAction.EXPORT_GENERATED,
        entreprise_id=document.entreprise_id, resource_type="document",
        resource_id=document.id,
        description=f"Export {normalized_format.upper()} des données extraites.",
        metadata={"format": normalized_format},
    )
    db.commit()

    return Response(
        content=content,

        media_type=media_type,

        headers={
            "Content-Disposition": (
                "attachment; filename=\""
                f"{filename}"
                "\""
            )
        },
    )


@router.patch("/{document_id}/entreprise", response_model=DocumentDetailOut)
def attribuer_document_entreprise(
    document_id: uuid.UUID,
    payload: DocumentEntrepriseUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Attribue manuellement un document puis relance son traitement sûr."""
    document = _get_document_or_404(db, document_id, current_user.cabinet_id)
    workflow_comptable_service.verifier_document_modifiable(db, document)
    entreprise = db.execute(select(Entreprise).where(
        Entreprise.id == payload.entreprise_id,
        Entreprise.cabinet_id == current_user.cabinet_id,
        Entreprise.is_active.is_(True),
        Entreprise.creee_automatiquement.is_(False),
    )).scalar_one_or_none()
    if entreprise is None:
        raise HTTPException(status_code=404, detail="Entreprise confirmée introuvable dans votre cabinet.")

    ancien_entreprise_id = document.entreprise_id
    _delete_accounting_data(db, document.id)
    donnees = dict(document.donnees_extraites or {})
    donnees["_entreprise_forcee_id"] = str(entreprise.id)
    donnees["source_identification_entreprise"] = "attribution_manuelle"
    document.donnees_extraites = donnees
    document.entreprise_id = entreprise.id
    document.statut = StatutDocumentEnum.EN_ATTENTE
    document.message_erreur = None
    document.type_erreur = None
    document.error_code = None
    document.saisie_topaze = False
    audit_service.enregistrer(
        db,
        user=current_user,
        action=audit_service.AuditAction.DOCUMENT_UPDATED,
        entreprise_id=entreprise.id,
        resource_type="document",
        resource_id=document.id,
        description="Attribution manuelle du document à une entreprise confirmée.",
        avant={"entreprise_id": ancien_entreprise_id},
        apres={"entreprise_id": entreprise.id, "retraitement": True},
    )
    db.commit()
    db.refresh(document)
    process_document.delay(str(document.id))
    return _build_document_detail(db, document)


@router.post(
    "/{document_id}/retraiter",
    response_model=DocumentOut,
)
def retraiter_document(
    document_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    document = _get_document_or_404(
        db,
        document_id,
        current_user.cabinet_id,
    )
    workflow_comptable_service.verifier_document_modifiable(db, document)

    file_path = Path(
        document.chemin_stockage
    )

    if not file_path.exists():
        raise HTTPException(
            status_code=404,
            detail=(
                "Fichier original introuvable. "
                "Retraitement impossible."
            ),
        )

    avant_retraitement = {
        "statut": document.statut,
        "categorie": document.categorie,
        "entreprise_id": document.entreprise_id,
        "saisie_topaze": document.saisie_topaze,
    }
    entreprise_id_audit = document.entreprise_id

    _delete_accounting_data(
        db,
        document.id,
    )

    document.statut = (
        StatutDocumentEnum.EN_ATTENTE
    )

    document.texte_ocr = None
    document.donnees_extraites = None
    document.message_erreur = None
    document.type_erreur = None
    document.error_code = None

    document.categorie = None
    document.annee = None
    document.mois = None
    document.entreprise_id = None

    document.saisie_topaze = False

    audit_service.enregistrer(
        db, user=current_user,
        action=audit_service.AuditAction.DOCUMENT_REPROCESSED,
        entreprise_id=entreprise_id_audit, resource_type="document",
        resource_id=document.id,
        description="Relance du traitement OCR/IA du document.",
        avant=avant_retraitement,
        apres={"statut": StatutDocumentEnum.EN_ATTENTE, "saisie_topaze": False},
    )
    db.commit()
    db.refresh(document)

    process_document.delay(
        str(document.id)
    )

    return document


@router.delete(
    "/{document_id}"
)
def delete_document(
    document_id: uuid.UUID,

    supprimer_fichier: bool = Query(
        default=True
    ),

    db: Session = Depends(get_db),

    current_user: User = Depends(
        get_current_user
    ),
):
    """
    Supprime définitivement:
    - le document de la base;
    - ses écritures comptables;
    - ses mouvements bancaires;
    - le fichier physique, lorsqu'il est présent dans STORAGE_PATH.

    Une fois la transaction validée, le hash du document disparaît de
    la base et le même fichier peut être importé à nouveau.
    """
    document = _get_document_or_404(
        db,
        document_id,
        current_user.cabinet_id,
    )
    workflow_comptable_service.verifier_document_modifiable(db, document)

    physical_path = document.chemin_stockage

    try:
        _delete_accounting_data(
            db,
            document.id,
        )

        audit_service.enregistrer(
            db, user=current_user, action=audit_service.AuditAction.DOCUMENT_DELETED,
            entreprise_id=document.entreprise_id,
            resource_type="document", resource_id=document.id,
            description=f"Suppression contrôlée du document {document.nom_fichier_original}.",
            avant={"nom_fichier": document.nom_fichier_original}, apres=None,
        )
        db.delete(document)
        db.commit()

    except Exception as exc:
        db.rollback()
        logger.exception(
            "Échec de la suppression du document %s",
            document_id,
        )
        raise HTTPException(
            status_code=500,
            detail="La suppression en base de données a échoué.",
        ) from exc

    file_deleted = False

    if supprimer_fichier:
        file_deleted = _delete_physical_file(
            physical_path
        )

        if not file_deleted:
            logger.warning(
                "Document %s supprimé de la base, mais le fichier "
                "physique n'a pas pu être supprimé: %s",
                document_id,
                physical_path,
            )

    return {
        "message": "Document supprimé définitivement.",
        "document_id": str(document_id),
        "fichier_supprime": file_deleted,
    }
