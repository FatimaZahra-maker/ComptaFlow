"""
app/services/export_service.py

Génère les exports d'un registre comptable (mêmes données que
GET /accounting/registers, Phase 5) en 3 formats génériques : Excel
(.xlsx), CSV, PDF. Génère aussi un export Topaze par écriture unique.
Ne stocke rien sur disque de façon permanente : les fichiers sont
générés en mémoire puis renvoyés directement au client (voir
app/api/export.py).

Topaze n'a pas de format d'import officiel documenté publiquement --
generer_export_topaze() produit un CSV pivot générique (date;compte;
libellé;débit;crédit), lisible en import manuel par la plupart des
logiciels comptables marocains. Si Topaze exige des colonnes/encodage
précis, fournir sa doc d'import pour ajuster cette fonction exactement.
"""
import csv
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

from app.models.ecriture import EcritureComptable

_ENTETES = ["Date", "Tiers", "N° pièce", "HT (MAD)", "TVA (MAD)", "TTC (MAD)"]


def _ligne_registre(ecriture: EcritureComptable) -> list[str]:
    return [
        ecriture.date_piece.strftime("%d/%m/%Y") if ecriture.date_piece else "—",
        ecriture.tiers or "—",
        ecriture.numero_piece or "—",
        f"{ecriture.montant_ht:.2f}",
        f"{ecriture.montant_tva:.2f}",
        f"{ecriture.montant_ttc:.2f}",
    ]


def generer_excel(lignes: list[EcritureComptable], titre: str, totaux: dict[str, Decimal]) -> bytes:
    """Génère un classeur Excel avec le détail des écritures + une ligne de totaux."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Registre"

    feuille.append([titre])
    feuille["A1"].font = Font(bold=True, size=14)
    feuille.append([])  # ligne vide
    feuille.append(_ENTETES)
    for cellule in feuille[3]:
        cellule.font = Font(bold=True)

    for ecriture in lignes:
        feuille.append(_ligne_registre(ecriture))

    feuille.append([])
    feuille.append([
        "TOTAL", "", "",
        f"{totaux['total_ht']:.2f}", f"{totaux['total_tva']:.2f}", f"{totaux['total_ttc']:.2f}",
    ])
    for cellule in feuille[feuille.max_row]:
        cellule.font = Font(bold=True)

    for index, largeur in enumerate([14, 25, 16, 14, 14, 14], start=1):
        feuille.column_dimensions[feuille.cell(row=3, column=index).column_letter].width = largeur

    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon.read()


def generer_csv(lignes: list[EcritureComptable], totaux: dict[str, Decimal]) -> bytes:
    """Génère un CSV simple, encodage UTF-8 avec BOM pour une ouverture
    correcte dans Excel (évite les accents mal affichés)."""
    tampon = io.StringIO()
    writer = csv.writer(tampon, delimiter=";")
    writer.writerow(_ENTETES)
    for ecriture in lignes:
        writer.writerow(_ligne_registre(ecriture))
    writer.writerow([])
    writer.writerow(["TOTAL", "", "", f"{totaux['total_ht']:.2f}", f"{totaux['total_tva']:.2f}", f"{totaux['total_ttc']:.2f}"])

    return ("\ufeff" + tampon.getvalue()).encode("utf-8")


def generer_pdf(lignes: list[EcritureComptable], titre: str, totaux: dict[str, Decimal]) -> bytes:
    """Génère un PDF tabulaire simple du registre, avec ligne de totaux en gras."""
    tampon = io.BytesIO()
    document = SimpleDocTemplate(tampon, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    elements = [Paragraph(titre, styles["Title"]), Spacer(1, 12)]

    donnees = [_ENTETES]
    for ecriture in lignes:
        donnees.append(_ligne_registre(ecriture))
    donnees.append([
        "TOTAL", "", "",
        f"{totaux['total_ht']:.2f}", f"{totaux['total_tva']:.2f}", f"{totaux['total_ttc']:.2f}",
    ])

    tableau = Table(donnees, repeatRows=1)
    tableau.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e40af")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
    ]))
    elements.append(tableau)
    document.build(elements)

    tampon.seek(0)
    return tampon.read()


def generer_export_topaze(ecriture: EcritureComptable) -> bytes:
    """
    Export CSV pivot pour une écriture unique (bouton "Exporter Topaze"
    du frontend, DocumentDetailPage.tsx) : date;compte;libellé;débit;
    crédit. Schéma comptable marocain standard, direction déterminée
    par ecriture.type_ecriture (achat vs vente).
    """
    tampon = io.StringIO()
    writer = csv.writer(tampon, delimiter=";")
    writer.writerow(["Date", "Compte", "Libellé", "Débit", "Crédit"])

    date_str = ecriture.date_piece.strftime("%d/%m/%Y") if ecriture.date_piece else ""
    tiers = ecriture.tiers or "Tiers"
    type_str = ecriture.type_ecriture.value if hasattr(ecriture.type_ecriture, "value") else str(ecriture.type_ecriture)
    est_vente = "vente" in type_str.lower()

    if est_vente:
        writer.writerow([date_str, "3421", f"Client {tiers}", f"{ecriture.montant_ttc:.2f}", ""])
        writer.writerow([date_str, "7111", "Ventes de marchandises", "", f"{ecriture.montant_ht:.2f}"])
        writer.writerow([date_str, "4455", "État — TVA facturée", "", f"{ecriture.montant_tva:.2f}"])
    else:
        writer.writerow([date_str, "6111", "Achats marchandises", f"{ecriture.montant_ht:.2f}", ""])
        writer.writerow([date_str, "34552", "TVA déductible sur achats", f"{ecriture.montant_tva:.2f}", ""])
        writer.writerow([date_str, "4411", f"Fournisseur {tiers}", "", f"{ecriture.montant_ttc:.2f}"])

    return ("\ufeff" + tampon.getvalue()).encode("utf-8")